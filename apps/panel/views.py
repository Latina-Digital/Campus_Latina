import csv
import io
from datetime import datetime, timedelta

from django.contrib import messages
from django.contrib.auth.decorators import user_passes_test
from django.contrib.auth.mixins import LoginRequiredMixin, UserPassesTestMixin
from django.db import transaction
from django.db.models import Avg, Count, Q
from django.db.models.functions import TruncMonth
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse, reverse_lazy
from django.utils import timezone
from django.utils.html import format_html
from django.views.decorators.http import require_POST
from django.views.generic import CreateView, DeleteView, ListView, UpdateView

from apps.accounts.models import User
from apps.assessments.models import Attempt, Choice, Question, Quiz
from apps.assessments.services import attempt_questions
from apps.catalog.models import Product, ProductLine
from apps.certificates.models import Certificate
from apps.core.models import Announcement
from apps.courses.models import Category, Course, Enrollment, LearningPath, Lesson, LessonProgress, Module, Resource

from . import forms
from .importers import read_questions, save_questions


def staff_required(view):
    return user_passes_test(lambda u: u.is_authenticated and u.is_staff, login_url="accounts:login")(view)


def admin_required(view):
    return user_passes_test(lambda u: u.is_authenticated and u.is_admin, login_url="accounts:login")(view)


class StaffMixin(LoginRequiredMixin, UserPassesTestMixin):
    section = ""
    admin_only = False

    def test_func(self):
        user = self.request.user
        return user.is_admin if self.admin_only else user.is_staff

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx["section"] = self.section
        return ctx


def badge(text, tone="neutral"):
    return format_html('<span class="pill pill-{}">{}</span>', tone, text)


def yes_no(value, yes="Sí", no="No"):
    return badge(yes, "success") if value else badge(no, "neutral")


# --------------------------------------------------------------------------- CRUD genérico
class PanelList(StaffMixin, ListView):
    template_name = "panel/generic_list.html"
    paginate_by = 25
    title = ""
    subtitle = ""
    columns = []  # [(encabezado, función(obj))]
    search_fields = []
    create_url = ""
    edit_url = ""
    delete_url = ""
    extra_actions = []  # [(label, icon, url_name)]

    def get_queryset(self):
        qs = super().get_queryset()
        q = self.request.GET.get("q", "").strip()
        if q and self.search_fields:
            cond = Q()
            for f in self.search_fields:
                cond |= Q(**{f"{f}__icontains": q})
            qs = qs.filter(cond)
        return qs

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx.update(
            title=self.title,
            subtitle=self.subtitle,
            headers=[c[0] for c in self.columns],
            rows=[
                {
                    "cells": [c[1](obj) for c in self.columns],
                    "edit": reverse(self.edit_url, args=[obj.pk]) if self.edit_url else None,
                    "delete": reverse(self.delete_url, args=[obj.pk]) if self.delete_url else None,
                    "extra": [(label, icn, reverse(name, args=[obj.pk])) for label, icn, name in self.extra_actions],
                }
                for obj in ctx["object_list"]
            ],
            create_url=reverse(self.create_url) if self.create_url else None,
            searchable=bool(self.search_fields),
            q=self.request.GET.get("q", ""),
        )
        return ctx


class PanelFormMixin(StaffMixin):
    template_name = "panel/generic_form.html"
    title = ""
    list_url = ""

    def get_success_url(self):
        return reverse(self.list_url)

    def form_valid(self, form):
        messages.success(self.request, "Cambios guardados.")
        return super().form_valid(form)

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx["title"] = self.title if not self.object else f"Editar: {self.object}"
        ctx["back_url"] = reverse(self.list_url)
        return ctx


class PanelCreate(PanelFormMixin, CreateView):
    object = None


class PanelUpdate(PanelFormMixin, UpdateView):
    pass


class PanelDelete(StaffMixin, DeleteView):
    template_name = "panel/confirm_delete.html"
    list_url = ""

    def get_success_url(self):
        return reverse(self.list_url)

    def form_valid(self, form):
        messages.success(self.request, f"«{self.object}» fue eliminado.")
        return super().form_valid(form)

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx["back_url"] = reverse(self.list_url)
        return ctx


def crud(model, form_class, prefix, section, title, title_plural, columns, search=(), admin_only=False, ordering=None, **list_extra):
    """Crea las 4 vistas CRUD de un modelo con una sola declaración."""
    attrs = {"section": section, "admin_only": admin_only}
    qs = model.objects.all() if ordering is None else model.objects.order_by(*ordering)
    lst = type(f"{model.__name__}List", (PanelList,), {
        **attrs, "model": model, "queryset": qs, "title": title_plural, "columns": columns, "search_fields": list(search),
        "create_url": f"panel:{prefix}_create", "edit_url": f"panel:{prefix}_edit", "delete_url": f"panel:{prefix}_delete", **list_extra,
    })
    create = type(f"{model.__name__}Create", (PanelCreate,), {**attrs, "model": model, "form_class": form_class, "title": f"Nuevo: {title}", "list_url": f"panel:{prefix}_list"})
    update = type(f"{model.__name__}Update", (PanelUpdate,), {**attrs, "model": model, "form_class": form_class, "list_url": f"panel:{prefix}_list"})
    delete = type(f"{model.__name__}Delete", (PanelDelete,), {**attrs, "model": model, "list_url": f"panel:{prefix}_list"})
    return lst.as_view(), create.as_view(), update.as_view(), delete.as_view()


category_views = crud(
    Category, forms.CategoryForm, "category", "categories", "categoría", "Categorías de cursos",
    [("Nombre", lambda o: format_html('<i data-lucide="{}" class="icon"></i> {}', o.icon, o.name)),
     ("Cursos", lambda o: o.courses.count()), ("Orden", lambda o: o.order)],
    search=("name",),
)
line_views = crud(
    ProductLine, forms.ProductLineForm, "line", "lines", "línea de productos", "Líneas de productos",
    [("Nombre", lambda o: format_html('<i data-lucide="{}" class="icon"></i> {}', o.icon, o.name)),
     ("Productos", lambda o: o.products.count()), ("Orden", lambda o: o.order)],
    search=("name",),
)
product_views = crud(
    Product, forms.ProductForm, "product", "products", "producto", "Productos del catálogo",
    [("Producto", lambda o: format_html('<div class="cell-media">{}<strong>{}</strong></div>',
                                        format_html('<img src="{}" alt="">', o.image.url) if o.image else "", o.name)),
     ("Marca", lambda o: badge(o.get_brand_display(), "navy" if o.brand == "latina" else "red")),
     ("Línea", lambda o: o.line or "—"), ("Activo", lambda o: yes_no(o.is_active))],
    search=("name", "summary"), ordering=("brand", "order", "name"),
)
announcement_views = crud(
    Announcement, forms.AnnouncementForm, "announcement", "announcements", "anuncio", "Anuncios",
    [("Título", lambda o: o.title), ("Tipo", lambda o: badge(o.get_tone_display(), {"urgente": "red", "nuevo": "gold"}.get(o.tone, "navy"))),
     ("Visible", lambda o: yes_no(o.is_active)), ("Fecha", lambda o: timezone.localtime(o.created_at).strftime("%d/%m/%Y"))],
    search=("title", "body"),
)


# --------------------------------------------------------------------------- Dashboard
@staff_required
def dashboard(request):
    now = timezone.now()
    month_start = timezone.localtime().replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    finished = Attempt.objects.filter(finished_at__isnull=False)
    n_finished = finished.count()
    six_months = (month_start - timedelta(days=160)).replace(day=1)

    def monthly(qs, field):
        rows = qs.filter(**{f"{field}__gte": six_months}).annotate(m=TruncMonth(field)).values("m").annotate(n=Count("id"))
        return {r["m"].strftime("%Y-%m"): r["n"] for r in rows if r["m"]}

    months = []
    cursor = six_months
    while cursor <= month_start:
        months.append(cursor)
        cursor = (cursor + timedelta(days=32)).replace(day=1)
    certs_m = monthly(Certificate.objects.filter(is_revoked=False), "issued_at")
    attempts_m = monthly(finished, "finished_at")
    month_names = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]

    course_scores = (
        Course.objects.filter(quiz__attempts__finished_at__isnull=False)
        .annotate(avg=Avg("quiz__attempts__score"), n=Count("quiz__attempts"))
        .order_by("-n")[:8]
    )
    brand_counts = dict(Enrollment.objects.values("course__brand").annotate(n=Count("id")).values_list("course__brand", "n"))
    active_users = User.objects.filter(is_active=True, is_superuser=False)
    learners_with_activity = LessonProgress.objects.values("user").distinct().count()

    context = {
        "section": "dashboard",
        "kpis": {
            "users": active_users.count(),
            "users_new": active_users.filter(date_joined__gte=month_start).count(),
            "active_learners": learners_with_activity,
            "courses": Course.objects.filter(is_published=True).count(),
            "drafts": Course.objects.filter(is_published=False).count(),
            "certificates": Certificate.objects.filter(is_revoked=False).count(),
            "certificates_month": Certificate.objects.filter(is_revoked=False, issued_at__gte=month_start).count(),
            "approval": round(finished.filter(passed=True).count() * 100 / n_finished) if n_finished else 0,
            "avg_score": round(finished.aggregate(a=Avg("score"))["a"] or 0),
            "attempts": n_finished,
        },
        "chart_activity": {
            "labels": [f"{month_names[m.month - 1]} {m:%y}" for m in months],
            "certificates": [certs_m.get(m.strftime("%Y-%m"), 0) for m in months],
            "attempts": [attempts_m.get(m.strftime("%Y-%m"), 0) for m in months],
        },
        "chart_courses": {
            "labels": [c.title[:28] for c in course_scores],
            "scores": [round(c.avg or 0) for c in course_scores],
        },
        "chart_brands": {
            "labels": ["Industrial Latina", "Induless", "Corporativo"],
            "values": [brand_counts.get("latina", 0), brand_counts.get("induless", 0), brand_counts.get("corporativo", 0)],
        },
        "recent_attempts": finished.select_related("user", "quiz__course").order_by("-finished_at")[:8],
        "top_courses": Course.objects.annotate(n=Count("enrollments")).order_by("-n")[:5],
        "inactive_users": active_users.exclude(pk__in=LessonProgress.objects.values("user")).count(),
        "now": now,
    }
    return render(request, "panel/dashboard.html", context)


# --------------------------------------------------------------------------- Usuarios
class UserList(PanelList):
    section = "users"
    admin_only = True
    model = User
    title = "Usuarios del campus"
    template_name = "panel/users.html"
    search_fields = ["username", "first_name", "last_name", "supervisor", "email"]
    subtitle = "Personas que ya ingresaron al campus. Los usuarios, roles y supervisores se administran en Syspro."
    columns = [
        ("Nombre", lambda u: format_html('<div class="cell-user"><span class="avatar avatar-initials" style="width:32px;height:32px;font-size:12px">{}</span><div><strong>{}</strong><small>@{}</small></div></div>', u.initials, u.display_name, u.username)),
        ("Rol", lambda u: badge(u.get_role_display(), "gold" if u.is_staff else "navy")),
        ("Supervisor", lambda u: u.supervisor or "—"),
        ("Origen", lambda u: u.get_origin_display()),
        ("Estado", lambda u: yes_no(u.is_active, "Activo", "Inactivo")),
        ("Último ingreso", lambda u: timezone.localtime(u.last_login).strftime("%d/%m/%Y %H:%M") if u.last_login else "Nunca"),
    ]

    def get_queryset(self):
        qs = super().get_queryset()
        role = self.request.GET.get("rol")
        return qs.filter(role=role) if role else qs

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx["roles"] = User.Role.choices
        ctx["role"] = self.request.GET.get("rol", "")
        return ctx


# --------------------------------------------------------------------------- Cursos
class CourseList(PanelList):
    section = "courses"
    model = Course
    title = "Cursos"
    template_name = "panel/courses.html"
    search_fields = ["title", "summary"]
    create_url = "panel:course_create"

    def get_queryset(self):
        return super().get_queryset().select_related("category").annotate(
            n_lessons=Count("modules__lessons", distinct=True), n_enrolled=Count("enrollments", distinct=True),
            n_certs=Count("certificates", distinct=True),
        ).order_by("order", "-created_at")


class CourseCreate(PanelCreate):
    section = "courses"
    model = Course
    form_class = forms.CourseForm
    title = "Nuevo curso"
    list_url = "panel:course_list"

    def get_success_url(self):
        return reverse("panel:course_builder", args=[self.object.pk])


class CourseUpdate(PanelUpdate):
    section = "courses"
    model = Course
    form_class = forms.CourseForm
    list_url = "panel:course_list"

    def get_success_url(self):
        return reverse("panel:course_builder", args=[self.object.pk])


class CourseDelete(PanelDelete):
    section = "courses"
    model = Course
    list_url = "panel:course_list"


@staff_required
def course_builder(request, pk):
    course = get_object_or_404(Course, pk=pk)
    quiz = course.quiz_or_none
    module_form = forms.ModuleForm(prefix="module", initial={"order": course.modules.count() + 1})
    resource_form = forms.ResourceForm(prefix="resource")
    quiz_form = forms.QuizForm(prefix="quiz", instance=quiz or Quiz(course=course))
    if request.method == "POST":
        action = request.POST.get("action")
        if action == "add_module":
            module_form = forms.ModuleForm(request.POST, prefix="module")
            if module_form.is_valid():
                module = module_form.save(commit=False)
                module.course = course
                module.save()
                messages.success(request, f"Módulo «{module}» creado.")
                return redirect("panel:course_builder", pk=pk)
        elif action == "add_resource":
            resource_form = forms.ResourceForm(request.POST, request.FILES, prefix="resource")
            if resource_form.is_valid():
                resource = resource_form.save(commit=False)
                resource.course = course
                resource.save()
                messages.success(request, "Material agregado.")
                return redirect("panel:course_builder", pk=pk)
        elif action == "save_quiz":
            quiz_form = forms.QuizForm(request.POST, prefix="quiz", instance=quiz or Quiz(course=course))
            if quiz_form.is_valid():
                quiz_form.save()
                messages.success(request, "Configuración de la evaluación guardada.")
                return redirect("panel:course_builder", pk=pk)
        elif action == "toggle_publish":
            course.is_published = not course.is_published
            course.save(update_fields=["is_published", "updated_at"])
            messages.success(request, "Curso publicado." if course.is_published else "Curso pasado a borrador.")
            return redirect("panel:course_builder", pk=pk)
        elif action == "toggle_coming_soon":
            course.coming_soon = not course.coming_soon
            course.save(update_fields=["coming_soon", "updated_at"])
            messages.success(
                request,
                "Curso marcado como Próximamente: se ve en el catálogo, pero nadie puede tomarlo."
                if course.coming_soon else "Curso habilitado: ya se puede tomar.",
            )
            return redirect("panel:course_builder", pk=pk)
    context = {
        "section": "courses",
        "course": course,
        "modules": course.modules.prefetch_related("lessons"),
        "module_form": module_form,
        "resource_form": resource_form,
        "quiz_form": quiz_form,
        "quiz": quiz,
        "question_count": quiz.questions.count() if quiz else 0,
        "stats": {
            "enrolled": course.enrollments.count(),
            "completed": course.enrollments.filter(completed_at__isnull=False).count(),
            "certificates": course.certificates.count(),
            "avg": round(Attempt.objects.filter(quiz__course=course, finished_at__isnull=False).aggregate(a=Avg("score"))["a"] or 0),
        },
    }
    return render(request, "panel/course_builder.html", context)


@staff_required
def module_edit(request, pk):
    module = get_object_or_404(Module, pk=pk)
    form = forms.ModuleForm(request.POST or None, instance=module)
    if request.method == "POST":
        if request.POST.get("action") == "delete":
            course_pk = module.course_id
            module.delete()
            messages.success(request, "Módulo eliminado.")
            return redirect("panel:course_builder", pk=course_pk)
        if form.is_valid():
            form.save()
            messages.success(request, "Módulo actualizado.")
            return redirect("panel:course_builder", pk=module.course_id)
    return render(request, "panel/generic_form.html", {
        "section": "courses", "form": form, "title": f"Editar módulo: {module}",
        "back_url": reverse("panel:course_builder", args=[module.course_id]), "deletable": True,
    })


@staff_required
def lesson_form(request, course_pk, pk=None):
    course = get_object_or_404(Course, pk=course_pk)
    lesson = get_object_or_404(Lesson, pk=pk, module__course=course) if pk else None
    if not course.modules.exists():
        messages.warning(request, "Primero crea un módulo para el curso.")
        return redirect("panel:course_builder", pk=course.pk)
    initial = {}
    if not lesson:
        module_id = request.GET.get("modulo")
        module = course.modules.filter(pk=module_id).first() or course.modules.last()
        initial = {"module": module, "order": module.lessons.count() + 1}
    form = forms.LessonForm(request.POST or None, request.FILES or None, instance=lesson, course=course, initial=initial)
    if request.method == "POST":
        if lesson and request.POST.get("action") == "delete":
            lesson.delete()
            messages.success(request, "Lección eliminada.")
            return redirect("panel:course_builder", pk=course.pk)
        if form.is_valid():
            form.save()
            messages.success(request, "Lección guardada.")
            if request.POST.get("action") == "save_add":
                return redirect(reverse("panel:lesson_create", args=[course.pk]) + f"?modulo={form.instance.module_id}")
            return redirect("panel:course_builder", pk=course.pk)
    return render(request, "panel/lesson_form.html", {"section": "courses", "form": form, "course": course, "lesson": lesson})


@staff_required
@require_POST
def reorder(request, kind, pk):
    model = {"module": Module, "lesson": Lesson}[kind]
    obj = get_object_or_404(model, pk=pk)
    siblings = list(obj.course.modules.all() if kind == "module" else obj.module.lessons.all())
    idx = siblings.index(obj)
    swap = idx - 1 if request.POST.get("dir") == "up" else idx + 1
    if 0 <= swap < len(siblings):
        siblings[idx], siblings[swap] = siblings[swap], siblings[idx]
        with transaction.atomic():
            for order, item in enumerate(siblings, start=1):
                if item.order != order:
                    item.order = order
                    item.save(update_fields=["order"])
    course_pk = obj.course_id if kind == "module" else obj.module.course_id
    return redirect(reverse("panel:course_builder", args=[course_pk]) + f"#{kind}-{pk}")


@staff_required
@require_POST
def resource_delete(request, pk):
    resource = get_object_or_404(Resource, pk=pk)
    resource.delete()
    messages.success(request, "Material eliminado.")
    return redirect("panel:course_builder", pk=resource.course_id)


# --------------------------------------------------------------------------- Banco de preguntas
def _get_quiz(course):
    quiz, _ = Quiz.objects.get_or_create(course=course)
    return quiz


@staff_required
def questions(request, course_pk):
    course = get_object_or_404(Course, pk=course_pk)
    quiz = _get_quiz(course)
    form = forms.QuestionForm(request.POST or None, prefix="new")
    import_form = forms.QuestionImportForm(prefix="imp")
    if request.method == "POST":
        action = request.POST.get("action")
        if action == "add" and form.is_valid():
            _save_question(quiz, form)
            messages.success(request, "Pregunta agregada al banco.")
            return redirect(reverse("panel:questions", args=[course.pk]) + "#nueva")
        if action == "import":
            import_form = forms.QuestionImportForm(request.POST, request.FILES, prefix="imp")
            if import_form.is_valid():
                try:
                    parsed, errors = read_questions(import_form.cleaned_data.get("file"), import_form.cleaned_data.get("text", ""))
                except Exception as exc:
                    parsed, errors = [], [f"No se pudo leer el archivo: {exc}"]
                if parsed:
                    n = save_questions(quiz, parsed, replace=import_form.cleaned_data.get("replace"))
                    messages.success(request, f"Se importaron {n} preguntas.")
                for error in errors[:10]:
                    messages.warning(request, error)
                if not parsed and not errors:
                    messages.warning(request, "No se encontraron preguntas con el formato esperado.")
                return redirect("panel:questions", course_pk=course.pk)
        if action == "delete_all":
            quiz.questions.all().delete()
            messages.success(request, "Se vació el banco de preguntas.")
            return redirect("panel:questions", course_pk=course.pk)
    context = {
        "section": "courses",
        "course": course,
        "quiz": quiz,
        "questions": quiz.questions.prefetch_related("choices"),
        "form": form,
        "import_form": import_form,
    }
    return render(request, "panel/questions.html", context)


@transaction.atomic
def _save_question(quiz, form, question=None):
    if question is None:
        question = Question(quiz=quiz, order=quiz.questions.count() + 1)
    question.text = form.cleaned_data["text"]
    question.explanation = form.cleaned_data["explanation"]
    question.save()
    question.choices.all().delete()
    for i, (text, ok) in enumerate(form.choices()):
        Choice.objects.create(question=question, text=text, is_correct=ok, order=i)
    return question


@staff_required
def question_edit(request, pk):
    question = get_object_or_404(Question.objects.select_related("quiz__course"), pk=pk)
    course = question.quiz.course
    choices = list(question.choices.all())
    initial = {"text": question.text, "explanation": question.explanation,
               "correct": next((i for i, c in enumerate(choices) if c.is_correct), 0)}
    initial.update({f"choice_{i}": c.text for i, c in enumerate(choices)})
    form = forms.QuestionForm(request.POST or None, initial=initial)
    if request.method == "POST":
        if request.POST.get("action") == "delete":
            question.delete()
            messages.success(request, "Pregunta eliminada.")
            return redirect("panel:questions", course_pk=course.pk)
        if form.is_valid():
            _save_question(question.quiz, form, question)
            messages.success(request, "Pregunta actualizada.")
            return redirect(reverse("panel:questions", args=[course.pk]) + f"#q-{question.pk}")
    return render(request, "panel/question_form.html", {"section": "courses", "form": form, "course": course, "question": question})


@staff_required
def question_template(request):
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = 'attachment; filename="plantilla_preguntas.csv"'
    response.write("﻿")
    writer = csv.writer(response, delimiter=";")
    writer.writerow(["pregunta", "a", "b", "c", "d", "correcta", "explicacion"])
    writer.writerow(["¿Cuántos días debe curar un enlucido antes de pintar?", "5 a 10 días", "15 a 20 días", "28 a 31 días", "40 a 50 días", "C", "El enlucido necesita casi un mes para curar."])
    return response


# --------------------------------------------------------------------------- Rutas
class PathList(PanelList):
    section = "paths"
    model = LearningPath
    title = "Rutas de aprendizaje"
    search_fields = ["title"]
    create_url = "panel:path_create"
    edit_url = "panel:path_edit"
    delete_url = "panel:path_delete"
    columns = [
        ("Ruta", lambda o: format_html("<strong>{}</strong><br><small class='muted'>{}</small>", o.title, o.summary)),
        ("Cursos", lambda o: o.steps.count()),
        ("Publicada", lambda o: yes_no(o.is_published)),
    ]


@staff_required
def path_form(request, pk=None):
    path = get_object_or_404(LearningPath, pk=pk) if pk else None
    form = forms.PathForm(request.POST or None, request.FILES or None, instance=path)
    formset = forms.PathStepFormSet(request.POST or None, instance=path or LearningPath())
    if request.method == "POST" and form.is_valid() and formset.is_valid():
        with transaction.atomic():
            path = form.save()
            formset.instance = path
            formset.save()
        messages.success(request, "Ruta guardada.")
        return redirect("panel:path_list")
    return render(request, "panel/path_form.html", {"section": "paths", "form": form, "formset": formset, "path": path})


class PathDelete(PanelDelete):
    section = "paths"
    model = LearningPath
    list_url = "panel:path_list"


# --------------------------------------------------------------------------- Resultados
def _filtered_attempts(request):
    qs = Attempt.objects.filter(finished_at__isnull=False).select_related("user", "quiz__course").order_by("-finished_at")
    g = request.GET
    if g.get("curso"):
        qs = qs.filter(quiz__course_id=g["curso"])
    if g.get("estado") == "aprobado":
        qs = qs.filter(passed=True)
    elif g.get("estado") == "reprobado":
        qs = qs.filter(passed=False)
    if g.get("supervisor"):
        qs = qs.filter(user__supervisor__icontains=g["supervisor"])
    if g.get("q"):
        qs = qs.filter(Q(user__first_name__icontains=g["q"]) | Q(user__last_name__icontains=g["q"]) | Q(user__username__icontains=g["q"]))
    for key, lookup in (("desde", "finished_at__date__gte"), ("hasta", "finished_at__date__lte")):
        try:
            if g.get(key):
                qs = qs.filter(**{lookup: datetime.strptime(g[key], "%Y-%m-%d").date()})
        except ValueError:
            pass
    if g.get("primer") == "1":
        firsts = {}
        for a in qs.order_by("started_at"):
            firsts.setdefault((a.user_id, a.quiz_id), a.pk)
        qs = qs.filter(pk__in=firsts.values())
    return qs


@staff_required
def results(request):
    qs = _filtered_attempts(request)
    export = request.GET.get("exportar")
    if export in {"xlsx", "csv"}:
        return _export_attempts(qs, export)
    from django.core.paginator import Paginator

    page = Paginator(qs, 30).get_page(request.GET.get("page"))
    agg = qs.aggregate(avg=Avg("score"), n=Count("id"), ok=Count("id", filter=Q(passed=True)))
    context = {
        "section": "results",
        "page": page,
        "courses": Course.objects.filter(quiz__isnull=False).order_by("title"),
        "g": request.GET,
        "summary": {"n": agg["n"], "avg": round(agg["avg"] or 0), "rate": round(agg["ok"] * 100 / agg["n"]) if agg["n"] else 0},
    }
    return render(request, "panel/results.html", context)


def _export_attempts(qs, fmt):
    header = ["Usuario", "Nombre", "Supervisor", "Área", "Curso", "Inicio", "Fin", "Duración", "Aciertos", "Total", "Nota %", "Estado", "Tiempo agotado"]
    rows = [
        [
            a.user.username, a.user.display_name, a.user.supervisor, a.user.area, a.quiz.course.title,
            timezone.localtime(a.started_at).strftime("%d/%m/%Y %H:%M"), timezone.localtime(a.finished_at).strftime("%d/%m/%Y %H:%M"),
            a.duration_display, a.correct_count, a.total, a.score, a.status_label, "Sí" if a.timed_out else "No",
        ]
        for a in qs
    ]
    stamp = timezone.localtime().strftime("%Y%m%d_%H%M")
    if fmt == "csv":
        response = HttpResponse(content_type="text/csv; charset=utf-8")
        response["Content-Disposition"] = f'attachment; filename="resultados_{stamp}.csv"'
        response.write("﻿")
        writer = csv.writer(response, delimiter=";")
        writer.writerow(header)
        writer.writerows(rows)
        return response
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Resultados"
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0A1C40")
    for row in rows:
        ws.append(row)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(42, max(10, *(len(str(c.value or "")) + 2 for c in col)))
    buf = io.BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="resultados_{stamp}.xlsx"'
    return response


@staff_required
def attempt_report(request, pk):
    attempt = get_object_or_404(Attempt.objects.select_related("user", "quiz__course"), pk=pk, finished_at__isnull=False)
    return render(request, "panel/attempt_report.html", {
        "attempt": attempt, "items": attempt_questions(attempt, with_answers=True), "printed_at": timezone.now(),
    })


# --------------------------------------------------------------------------- Certificados
@staff_required
def certificates(request):
    qs = Certificate.objects.select_related("user", "course").order_by("-issued_at")
    q = request.GET.get("q", "").strip()
    if q:
        qs = qs.filter(Q(code__icontains=q) | Q(holder_name__icontains=q) | Q(course_title__icontains=q))
    from django.core.paginator import Paginator

    return render(request, "panel/certificates.html", {
        "section": "certificates", "page": Paginator(qs, 30).get_page(request.GET.get("page")), "q": q,
    })


@admin_required
@require_POST
def certificate_toggle(request, pk):
    certificate = get_object_or_404(Certificate, pk=pk)
    certificate.is_revoked = not certificate.is_revoked
    certificate.save(update_fields=["is_revoked"])
    messages.success(request, f"Certificado {certificate.code} {'revocado' if certificate.is_revoked else 'restablecido'}.")
    return redirect(request.META.get("HTTP_REFERER") or "panel:certificates")
